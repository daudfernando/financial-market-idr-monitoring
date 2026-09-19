"""Create five editable checkpoint slides from the verified project snapshot.

Optional authoring dependency: python-pptx==1.0.2 (not a pipeline dependency).
"""
from pathlib import Path
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor

ROOT = Path(__file__).resolve().parents[1]
SLIDES = [
 ('Harga global, perspektif rupiah', '01 / MASALAH BISNIS', [
  'Harga USD dan kurs USD/IDR berubah pada frekuensi berbeda.',
  'Analis membutuhkan nilai IDR dan penjelasan kontribusi harga vs kurs dalam satu dataset.',
  'Lingkup terukur: 4 saham, 7.788 bar menit, 20 tanggal JISDOR.',
  'Skenario treasury / market risk; quantity satu saham. Tidak ada klaim portofolio atau penghematan bisnis aktual.'
 ], '5 menit. Jelaskan masalah penggabungan waktu sumber dan kebutuhan audit. Sumber angka: data/dashboard_readiness.json, audit 17 September 2026.'),
 ('Dua jalur, satu warehouse', '02 / ARSITEKTUR', [
  'Batch: BI → Airflow → GCS raw → Spark → GCS harian → BigQuery.',
  'Replay: histori Yahoo → Kafka → Python consumer → GCS → BigQuery.',
  'dbt menggabungkan FX, menghitung fitur/sinyal, dan menguji kualitas.',
  'Docker Compose menjalankan layanan lokal. Dataflow tidak diperlukan; guideline memperbolehkan replay.'
 ], '7 menit. Kafka menerima event; Python memvalidasi saat event datang. GCS/BQ dimuat setelah replay selesai. JISDOR dijadwalkan 18 WIB saat Docker aktif. Tidak menyebut harga live atau warehouse setiap tick.'),
 ('Data siap dipakai di BigQuery', '03 / HASIL PIPELINE', [
  'fact_jisdor_daily: 20 tanggal valid, sampai 17 September 2026.',
  'mart_stock_monitoring: 7.788 bar; JPM, BAC, GS, MS; sesi 10–16 September.',
  'mart_stock_latest: 4 bar terbaru berdasarkan waktu sumber.',
  'Rekonsiliasi 7.788 Kafka → Python → GCS → BigQuery cocok; duplikasi/mismatch/FX kosong = 0.'
 ], '10 menit termasuk demo Airflow, GCS dan BigQuery. Project jcdeah-009; dataset daud_finalproject untuk fact/input, daud_finalproject_demo untuk mart. Replay run 20260917T123103782750Z.'),
 ('Sinyal yang dapat dijelaskan', '04 / ANALITIK DAN DASHBOARD', [
  'SMA3 menyilang naik SMA8 → BUY; menyilang turun → SELL; tanpa silang → HOLD.',
  'Riwayat kurang/gap/FX kosong → NO_SIGNAL. Sinyal bar tersedia setelah menit selesai.',
  'Hasil histori: BUY 530 · SELL 522 · HOLD 6.489 · NO_SIGNAL 247.',
  'Dashboard: grafik nilai IDR, grafik kurs, sinyal, dan return relatif tiga saham pembanding.'
 ], '9 menit termasuk demo dashboard. Latest keempat saham semuanya HOLD: jangan memaksa variasi. Lima fixture sintetis terpisah menguji rumus. Parameter simulasi belum dibacktest dan bukan rekomendasi investasi.'),
 ('Bukti kualitas dan kesiapan', '05 / CHECKPOINT', [
  '8 view dbt + 21 data tests + 3 unit tests: seluruh build sukses.',
  'Batch terjadwal 17 September sukses; alert kegagalan lokal tersedia di Mailpit.',
  'Gap sumber: GS 11 menit, MS 1 menit; tidak diisi dengan angka buatan.',
  'Teknis/data siap. Pengumpulan: tinjau slide, publikasikan link GitHub, dan latihan presentasi.'
 ], '4 menit + 5 menit Q&A. Bedakan replay dari live; warehouse snapshot dari streaming insertion. Jelaskan FX unknown availability memakai tanggal sebelumnya. Dataflow arsip opsional. HTML memenuhi dua chart; report Looker dapat dibuat pengguna dari mart.'),
]


def main():
    deck = Presentation()
    deck.slide_width, deck.slide_height = Inches(13.333), Inches(7.5)
    for index, (title, eyebrow, bullets, notes) in enumerate(SLIDES, 1):
        slide = deck.slides.add_slide(deck.slide_layouts[6])
        slide.background.fill.solid()
        slide.background.fill.fore_color.rgb = RGBColor.from_string('F2F5ED')
        def box(x, y, width, height, text, size, color='183F34', bold=False):
            frame = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(width), Inches(height)).text_frame
            frame.word_wrap = True
            p = frame.paragraphs[0]
            p.text = text
            p.font.size, p.font.bold, p.font.color.rgb = Pt(size), bold, RGBColor.from_string(color)
            p.font.name = 'Aptos'
        box(.65, .35, 12, .4, eyebrow, 13, '557B66', True)
        box(.65, 1, 12, 1, title, 34, bold=True)
        for i, bullet in enumerate(bullets):
            box(.7, 2.2 + i * 1.02, 11.9, .95, bullet, 22)
        box(.65, 6.85, 12, .35, f'JCDEAH-009  |  17 September 2026  |  REPLAY HISTORIS ASLI  |  {index}/5', 11, '557B66')
        slide.notes_slide.notes_text_frame.text = notes
    output = ROOT / 'docs/final_project_checkpoint.pptx'
    deck.save(output)
    assert len(Presentation(output).slides) == 5
    print(output)


if __name__ == '__main__':
    main()
